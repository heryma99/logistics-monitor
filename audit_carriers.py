# -*- coding: utf-8 -*-
"""audit_carriers.py — 合联/DPEX 账单内部一致性审计（无外部价格表时的自检引擎）
DPEX: 同 Dest×重量段 的隐含单价（Total÷计费重）中位数偏离 >25% → 异常
合联: 合计(RMB) vs (转运费+海运费+贴标费+提货费) 自洽校验
铁律: 账单只读，结果写 data/carriers_audit.json
"""
import sys, os, json, glob, statistics
sys.path.insert(0, r'D:\WB文件\2026-09-09-11-43-53\.secrets\pylibs')
import openpyxl
from collections import defaultdict

HERE = os.path.dirname(os.path.abspath(__file__))
def num(v):
    try: return float(v)
    except: return None

# ---------- DPEX 全序列 ----------
dpex_out = []
for f in sorted(glob.glob(os.path.join(HERE, "bills", "*830125*.xlsx")) + glob.glob(os.path.join(HERE, "bills", "*账单(2026-8-16至31*"))):
    try:
        wb = openpyxl.load_workbook(f, read_only=True, data_only=True)
        ws = wb["Sheet1"]
        rows = list(ws.iter_rows(values_only=True))
        hdr_i = next(i for i, r in enumerate(rows) if r[0] == "S/No")
        hdr = [str(c) if c else "" for c in rows[hdr_i]]
        i_dest = hdr.index("Dest"); i_dead = hdr.index("Dead Wt (gm)"); i_vol = hdr.index("Vol. Wt (gm)")
        i_tot = hdr.index("Total Amount"); i_frt = hdr.index("Freight"); i_sur = hdr.index("Surcharge")
        groups = defaultdict(list)
        n = 0
        for r in rows[hdr_i+1:]:
            if r[0] is None: continue
            n += 1
            dead, vol = num(r[i_dead]) or 0, num(r[i_vol]) or 0
            chg_kg = max(dead, vol) / 1000.0
            tot = num(r[i_tot]) or 0
            if chg_kg <= 0 or tot <= 0: continue
            dest = str(r[i_dest] or "-")
            br = "<1kg" if chg_kg < 1 else "1-2" if chg_kg < 2 else "2-5" if chg_kg < 5 else "5-10" if chg_kg < 10 else "10-20" if chg_kg < 20 else "20+"
            rate = tot / chg_kg
            groups[(dest, br)].append({"no": str(r[3] or ""), "kg": round(chg_kg, 2), "rate": rate,
                                       "freight": num(r[i_frt]), "sur": num(r[i_sur]), "total": tot})
        # 离群检测
        outliers = []
        for (dest, br), items in groups.items():
            if len(items) < 3: continue
            rates = [x["rate"] for x in items]
            med = statistics.median(rates)
            for x in items:
                if med > 0 and abs(x["rate"] - med) / med > 0.25:
                    outliers.append({"period": os.path.basename(f)[:28], "dest": dest, "br": br,
                                     "no": x["no"], "kg": x["kg"], "rate": round(x["rate"], 1),
                                     "median": round(med, 1), "total": x["total"]})
        dpex_out.append({"file": os.path.basename(f), "tickets": n, "groups": len(groups), "outliers": outliers})
        print(os.path.basename(f)[:34], "|", n, "票 | 分组", len(groups), "| 离群", len(outliers))
    except Exception as e:
        print("FAIL", os.path.basename(f), str(e)[:60])

# ---------- 合联自洽 ----------
hl_issues = []
try:
    wb = openpyxl.load_workbook(os.path.join(HERE, "bills", "香港合联8月份运费对账单.xlsx"), read_only=True, data_only=True)
    ws = wb["Worksheet"]
    rows = list(ws.iter_rows(values_only=True))
    hdr_i = next(i for i, r in enumerate(rows) if r[0] == "出货日期")
    for r in rows[hdr_i+1:]:
        if r[0] is None or r[2] is None: continue
        total = num(r[14]) or 0
        parts = sum(num(x) or 0 for x in (r[10], r[11], r[12], r[13]))
        if abs(total - parts) > 0.5:
            hl_issues.append({"no": str(r[2]), "total": total, "parts": round(parts, 2),
                              "diff": round(total - parts, 2)})
except Exception as e:
    print("合联 FAIL", str(e)[:60])
print("合联自洽问题:", len(hl_issues))

out = {"generated_at": __import__("datetime").datetime.now().strftime("%Y-%m-%d %H:%M:%S"),
       "dpex": dpex_out, "hl_issues": hl_issues}
json.dump(out, open(os.path.join(HERE, "data", "carriers_audit.json"), "w", encoding="utf-8"), ensure_ascii=False, indent=1)
tot_out = sum(len(d["outliers"]) for d in dpex_out) + len(hl_issues)
print(f"审计完成：DPEX {len(dpex_out)} 期 + 合联 1 期，内部一致性异常 {tot_out} 条")
